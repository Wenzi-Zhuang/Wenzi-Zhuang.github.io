"""Read small user-supplied documents for one chat turn without storing files."""

import io
import zipfile

from docx import Document
from pypdf import PdfReader


def extract_upload_text(content: bytes, filename: str) -> str:
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if suffix in {".txt", ".md", ".csv"}:
            text = content.decode("utf-8-sig")
        elif suffix == ".docx":
            if not content.startswith(b"PK"):
                raise ValueError("文件内容不是有效 DOCX")
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 8 * 1024 * 1024:
                    raise ValueError("文档解压后过大，无法处理")
            document = Document(io.BytesIO(content))
            text = "\n".join(p.text for p in document.paragraphs)
            for table in document.tables:
                text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        elif suffix == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise ValueError("文件内容不是有效 PDF")
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted or len(reader.pages) > 40:
                raise ValueError("加密文件或超过 40 页的 PDF 暂时无法处理")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            raise ValueError("仅支持 PDF、DOCX、TXT、MD、CSV 文档")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("文件无法读取，请检查格式或改用可复制文字的文档") from exc
    text = text.strip()
    if not text:
        raise ValueError("未能从文件提取文字，请改用可复制文字的文档")
    return text[:10000]
