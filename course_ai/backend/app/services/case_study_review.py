"""Conservative, full-document checks for the undergraduate case-paper workflow.

The checklist paraphrases section III of the 2022 DUFE writing specification.
An unobservable Word/print property is *never* reported as a pass.
"""

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


CASE_FILE_BYTES = 1024 * 1024
CASE_DAILY_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 32 * 1024 * 1024

# Every numbered paragraph in section III is represented. Long paragraphs are
# split where they impose distinct tests, so the UI never hides an unchecked item.
FORMAT_RULES = (
    ("3.1.1", "纸张为 A4（210 × 297 mm）"),
    ("3.1.2", "上下左右页边距均为 2 cm，装订线为 0"),
    ("3.2.1", "全文横排、从左到右、通栏编辑"),
    ("3.2.2", "使用规范简体汉字，避免不必要的繁体、异体及不规范汉字"),
    ("3.3.1", "封面文头固定文字、3 号宋体加粗、居中及空行"),
    ("3.3.2", "封面主标题 2 号黑体加粗、居中及空行"),
    ("3.3.3", "副标题小 2 号黑体加粗、破折号、居中且正副标题不超过两行"),
    ("3.3.4", "封面作者等项目逐行排列；项目名 3 号黑体、内容 3 号楷体，专业方向和导师职称完整，答辩日期及成绩留空"),
    ("3.3.5a", "中文摘要另起一节；标题、字号、对齐及段前后间距合规"),
    ("3.3.5b", "中文摘要内容小 4 号宋体、首行缩进 2 字符、固定 22 磅行距"),
    ("3.3.5c", "中文关键词标题及内容字体字号、数量和词间空格合规且不加标点"),
    ("3.3.6a", "英文摘要另起一节；ABSTRACT 标题字体、字号、加粗、居中及段间距合规"),
    ("3.3.6b", "英文摘要内容小 4 号 Times New Roman、首行缩进 2 字符、固定 22 磅行距且英文标点后有空格"),
    ("3.3.6c", "Key words 标题及关键词字号、字体、左对齐和分号分隔合规"),
    ("3.3.7a", "目录另起一节；标题、字体字号、居中及段间距合规"),
    ("3.3.7b", "目录列出一、二级标题及页码；仿宋正文、Times New Roman 页码、1.5 倍行距及页码右齐"),
    ("3.3.8a", "正文另起一节；论文主副标题字体字号、居中及段间距合规"),
    ("3.3.8b", "正文所有汉字小 4 号宋体；英文数字字号相应且符号适当"),
    ("3.3.8c", "正文标准字距、原则上固定 22 磅行距；表格公式按规定另行处理"),
    ("3.3.9a", "一至五级标题层次及编号依次为一、（一）1.（1）①，并连续编号"),
    ("3.3.9b", "一级标题 4 号黑体、独占行、段前后各 1 行、单倍行距且无末尾标点"),
    ("3.3.9c", "二级标题小 4 号宋体、独占行、段前后各 0.5 行、单倍行距且无末尾标点"),
    ("3.3.9d", "三级以下标题字体字号及接排规则、末尾标点符合要求"),
    ("3.3.9e", "同级标题原则上不少于三个；标题和正文首行缩进 2 字符且不悬挂缩进"),
    ("3.3.10", "注释使用按页重编的①②脚注；字体字号、文献格式和引文页码区间合规"),
    ("3.3.11", "如有附录，其标题、字号、段间距及内容缩进与行距合规"),
    ("3.3.12a", "主要参考文献标题、位置、字体字号、中文在前外文在后及统一连续编号合规"),
    ("3.3.12b", "各类中文和外文文献的作者、年份、标题、刊名、卷期、页码及标点符合规范"),
    ("3.3.12c", "外文作者姓名顺序及缩写、卷期信息和标题实词首字母大写合规"),
    ("3.3.13", "后记标题和正文的字号、字体、居中、段距及行距合规"),
    ("3.3.14", "封三上半页预留指导教师评语、签名和日期"),
    ("3.3.15", "封三下半页预留学院答辩机构意见、签名和日期"),
    ("3.4.1", "每个表的表号、标题和计量单位位于表上，5 号黑体加粗、单倍行距且连续编号"),
    ("3.4.2", "统计表和会计表线型、开放侧边框、表内字体字号、单倍行距与数字对齐合规"),
    ("3.4.3", "表中数据有出处说明，置于表下或标题脚注"),
    ("3.4.4", "表格尽量不跨页，正文先引导再出现表；图表阿拉伯数字连续编号且题名位置正确"),
    ("3.5.1", "图居中、上下环绕；图号图名在图下、5 号黑体且连续编号"),
    ("3.5.2", "图中文字字号、线条画布及正文引导语符合要求"),
    ("3.6.1", "公式首行缩进、5 号宋体、单倍行距、段前后间距和右侧编号合规"),
    ("3.6.2", "公式连续编号、说明顶格、长公式换行和对齐合规"),
    ("3.7.1", "数字使用汉字或阿拉伯数字的场景正确且全文一致；五位以上阿拉伯数字按空格分组"),
    ("3.8.1", "中英文标点区分，全文无明显误用或混用"),
    ("3.9.1", "计量单位使用法定计量单位"),
    ("3.10.1", "封面无页码；中英文摘要、目录、正文各自从 1 重新编号，页码居中置于页脚"),
    ("3.11.1", "封面无页眉；摘要、目录、附录、参考文献和后记各有对应页眉、分节及取消链接前节"),
    ("3.11.2", "正文奇数页页眉为论文题目，偶数页为学校届别文字，作用范围为本节"),
    ("3.11.3", "页眉使用文武线，5 号宋体加粗且居中"),
)

STRUCTURE_ITEMS = ("封面", "中文摘要", "英文摘要", "目录", "正文", "注释", "主要参考文献", "后记", "封底")

TOPIC_CRITERIA = (
    "研究问题能从指定案例情境提炼，不能是任何企业皆适用的空泛问题",
    "研究问题有学术价值，而非仅描述事实或常识",
    "说明研究对象、核心现象及结果、机制或条件之间的关系",
    "范围聚焦、可用一句话清晰表达，避免堆砌变量",
)

SELF_EVALUATION_ITEMS = (
    "12 研究问题与案例类型（含 12-1）", "13 案例选择依据", "14 研究意义",
    "15 理论依据、引用、假设及理论增量（含 15-1 至 15-3）",
    "16 研究与分析方法、变量测量、对照及结论（含 16-1 至 16-5）",
    "17 现实问题、解决方案及政策建议（含 17-1 至 17-2）",
    "18 参考文献篇数、级别及发表年份分布",
)


@dataclass
class Paper:
    document: object
    text: str
    paragraphs: list
    sha256: str


def read_word(content: bytes, filename: str) -> Paper:
    if not filename.lower().endswith(".docx"):
        raise ValueError("请提交 Word .docx 文档；旧版 .doc 暂不支持逐项格式检查")
    if len(content) > CASE_FILE_BYTES:
        raise ValueError("文件过大，无法处理：案例论文单个文档不能超过 1 MB")
    if not content.startswith(b"PK"):
        raise ValueError("文件不是有效的 Word .docx 文档")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > 2000 or sum(item.file_size for item in infos) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("Word 文档解压后过大，无法处理")
            if any(item.filename.startswith("/") or ".." in item.filename.split("/") for item in infos):
                raise ValueError("Word 文档结构不安全")
        document = Document(io.BytesIO(content))
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Word 文档无法读取，请检查文件是否损坏") from exc
    # XML order includes nested tables and Word text boxes, which python-docx's
    # top-level Document.paragraphs collection omits.
    paragraphs = [Paragraph(element, document) for element in document.element.body.xpath(".//w:p")]
    text = "\n".join(p.text for p in paragraphs if p.text.strip())
    if not text.strip():
        raise ValueError("Word 文档中没有可检查的文字")
    return Paper(document, text, paragraphs, hashlib.sha256(content).hexdigest())


def inspect_format(paper: Paper) -> dict:
    """Run complete OOXML traversal, then identify rules needing visual/human review."""
    results = {code: {"code": code, "requirement": title, "status": "复核", "evidence": "该项不能仅凭 Word 结构可靠判定，请人工核对页面呈现。"}
               for code, title in FORMAT_RULES}

    def record(code: str, ok: bool, evidence: str) -> None:
        results[code]["status"] = "通过" if ok else "需修改"
        results[code]["evidence"] = evidence

    sections = paper.document.sections
    paper_size_bad = [i + 1 for i, s in enumerate(sections)
                      if abs(s.page_width.cm - 21) > .1 or abs(s.page_height.cm - 29.7) > .1]
    record("3.1.1", not paper_size_bad, f"已检查 {len(sections)} 节；非 A4 节号：{paper_size_bad or '无'}。")
    margin_bad = [i + 1 for i, s in enumerate(sections)
                  if any(abs(getattr(s, side).cm - 2) > .08 for side in
                         ("top_margin", "bottom_margin", "left_margin", "right_margin"))
                  or (s.gutter and abs(s.gutter.cm) > .02)]
    record("3.1.2", not margin_bad, f"已检查 {len(sections)} 节；页边距或装订线不符的节号：{margin_bad or '无'}。")

    normalized = re.sub(r"\s+", "", paper.text)
    if "东北财经大学本科毕业论文" in normalized:
        results["3.3.1"]["evidence"] = "已在全文找到固定文头；字体、加粗、空行和封面位置仍需人工核对。"
    else:
        record("3.3.1", False, "未找到规定的固定文头。")
    headings = {"3.3.5a": r"^摘\s*要$", "3.3.6a": r"^ABSTRACT$", "3.3.7a": r"^目\s*录$",
                "3.3.12a": r"^主要参考文献$", "3.3.13": r"^后\s*记$"}
    for code, pattern in headings.items():
        matches = [p for p in paper.paragraphs if re.fullmatch(pattern, p.text.strip(), re.I)]
        if not matches:
            record(code, False, "全文未找到该部分的独立标题。")
        else:
            results[code]["evidence"] = f"全文找到 {len(matches)} 个标题；仍须核对字体、字号、间距和分节。"
    for code, label in (("3.3.5c", "关键词"), ("3.3.6c", "Key words")):
        if label.lower() not in paper.text.lower():
            record(code, False, f"全文未找到“{label}”。")
        else:
            results[code]["evidence"] = f"已找到“{label}”；数量和字体等仍须复核。"
    chinese_keywords = next((p.text for p in paper.paragraphs if p.text.strip().startswith("关键词")), "")
    if chinese_keywords:
        terms = [item for item in re.split(r"[\s;；,，、]+", re.split(r"[:：]", chinese_keywords, 1)[-1]) if item]
        if terms and not 3 <= len(terms) <= 7:
            record("3.3.5c", False, f"中文关键词疑似 {len(terms)} 个（要求 3–7 个）：{terms[:10]}。")
    english_keywords = next((p.text for p in paper.paragraphs if p.text.strip().lower().startswith("key words")), "")
    if english_keywords and ";" not in english_keywords and "；" not in english_keywords:
        results["3.3.6c"]["evidence"] = "未在英文关键词行找到英文分号；请核对关键词是否按要求分隔。"

    body = paper.document.element.body
    tracked = len(body.xpath(".//w:ins")) + len(body.xpath(".//w:del"))
    details = {"sections": len(sections), "paragraphs": len(paper.paragraphs),
               "tables": len(body.xpath(".//w:tbl")), "tracked_changes": tracked,
               "characters": len(paper.text), "document_sha256": paper.sha256}
    if len(sections) < 5:
        record("3.10.1", False, f"全文只有 {len(sections)} 节，无法分别实现封面、中文摘要、英文摘要、目录和正文的独立页码要求。")
        record("3.11.1", False, f"全文只有 {len(sections)} 节，无法为各部分设置不同页眉。")
    else:
        restarts = 0
        for section in sections:
            page_type = section._sectPr.find(qn("w:pgNumType"))
            if page_type is not None and page_type.get(qn("w:start")) == "1":
                restarts += 1
        details["page_number_restarts_at_1"] = restarts
        if restarts < 3:
            record("3.10.1", False, f"仅找到 {restarts} 处页码从 1 重新开始，摘要、目录和正文至少需要三处。")
        else:
            results["3.10.1"]["evidence"] = f"找到 {restarts} 处页码从 1 重新开始；仍需确认各节对应关系及页脚居中。"
        linked_headers = [i + 1 for i, section in enumerate(sections[1:], 1) if section.header.is_linked_to_previous]
        results["3.11.1"]["evidence"] = f"后续节链接前节页眉的节号：{linked_headers or '无'}；仍需核对具体页眉文字和分节位置。"
    if not paper.document.settings.odd_and_even_pages_header_footer:
        record("3.11.2", False, "文档未启用奇偶页不同页眉，不能满足正文奇偶页页眉要求。")
    else:
        results["3.11.2"]["evidence"] = "已启用奇偶页不同页眉；仍需核对两种页眉文字与作用范围。"
    cover_header = "".join(p.text for p in sections[0].header.paragraphs).strip()
    if cover_header and not sections[0].different_first_page_header_footer:
        record("3.11.1", False, "封面所在首节存在非空页眉，且未启用首页不同页眉。")
    details["text_boxes"] = len(body.xpath(".//w:txbxContent"))
    details["inline_images"] = len(body.xpath(".//w:drawing"))
    # Traverse every nonempty paragraph and every run. A style inherited from a
    # template may be opaque, so this is a diagnostic, not a blanket pass.
    wrong_fonts = []
    for index, paragraph in enumerate(paper.paragraphs, 1):
        for run in paragraph.runs:
            if not re.search(r"[\u4e00-\u9fff]", run.text):
                continue
            east_asia = run._element.rPr.rFonts.get(qn("w:eastAsia")) if run._element.rPr is not None and run._element.rPr.rFonts is not None else None
            if east_asia and east_asia not in {"宋体", "SimSun", "黑体", "SimHei", "楷体", "楷体_GB2312", "仿宋", "FangSong"}:
                wrong_fonts.append({"paragraph": index, "font": east_asia})
                if len(wrong_fonts) >= 12:
                    break
        if len(wrong_fonts) >= 12:
            break
    if wrong_fonts:
        results["3.3.8b"]["evidence"] = "发现可能不符的显式中文字体（前 12 处）：" + str(wrong_fonts)
    else:
        results["3.3.8b"]["evidence"] = f"扫描了全部 {len(paper.paragraphs)} 段；未发现明显异常的显式中文字体，继承样式与字号仍需复核。"

    caption_numbers = [int(m.group(1)) for p in paper.paragraphs if (m := re.match(r"^表\s*(\d+)", p.text.strip()))]
    if caption_numbers:
        if caption_numbers != list(range(1, len(caption_numbers) + 1)):
            record("3.4.1", False, f"表号没有从 1 连续编号：{caption_numbers[:30]}。")
        else:
            results["3.4.1"]["evidence"] = f"表号连续：{caption_numbers[:30]}；标题样式、位置和计量单位仍需人工复核。"
    figure_numbers = [int(m.group(1)) for p in paper.paragraphs if (m := re.match(r"^图\s*(\d+)", p.text.strip()))]
    if figure_numbers:
        if figure_numbers != list(range(1, len(figure_numbers) + 1)):
            record("3.5.1", False, f"图号没有从 1 连续编号：{figure_numbers[:30]}。")
        else:
            results["3.5.1"]["evidence"] = f"图号连续：{figure_numbers[:30]}；图片位置、环绕和标题样式仍需人工复核。"
    if tracked:
        details["warning"] = "文档含修订痕迹；清洁版提交前请接受或拒绝全部修订。"
    summary = {status: sum(item["status"] == status for item in results.values())
               for status in ("通过", "需修改", "复核")}
    return {"items": list(results.values()), "summary": summary, "details": details,
            "notice": "已遍历正文、表格、文本框中的可读取段落和全文分节；嵌入对象、精确分页、文武线及打印效果不能可靠自动确认，明确标为复核而非通过。"}


def structure_report(paper: Paper) -> list[dict]:
    text = paper.text
    patterns = {"封面": r"东北财经大学本科毕业论文", "中文摘要": r"摘\s*要", "英文摘要": r"ABSTRACT",
                "目录": r"目\s*录", "正文": r"一、|引言|绪论", "注释": r"[①②]|脚注",
                "主要参考文献": r"主要参考文献", "后记": r"后\s*记", "封底": r"指导教师评语|学院答辩机构意见"}
    return [{"name": name, "found": bool(re.search(patterns[name], text, re.I)),
             "note": "通过文字线索定位；具体位置和内容仍需复核。"} for name in STRUCTURE_ITEMS]
