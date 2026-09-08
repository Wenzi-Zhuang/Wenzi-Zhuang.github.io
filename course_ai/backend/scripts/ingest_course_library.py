import asyncio
import hashlib
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pptx import Presentation
from sqlalchemy import delete

from app.config import get_settings
from app.database import SessionLocal
from app.models import SourceChunk

ALLOWED_DIRS = {"00_course_rules": 100, "01_textbook": 80, "02_lectures": 90,
                "03_accounting_standards": 70, "04_cases": 60, "05_examples_entries": 60,
                "20_knowledge_system": 85}
EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md", ".csv"}


def extract(path: Path):
    if path.suffix.lower() == ".pdf":
        return [(f"第{i}页", page.extract_text() or "") for i, page in enumerate(PdfReader(path).pages, 1)]
    if path.suffix.lower() == ".docx":
        return [(None, "\n".join(p.text for p in Document(path).paragraphs))]
    if path.suffix.lower() == ".pptx":
        return [(f"第{i}页", "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text")))
                for i, slide in enumerate(Presentation(path).slides, 1)]
    return [(None, path.read_text(encoding="utf-8", errors="ignore"))]


def chunks(text: str, max_chars=2400):
    blocks = [x.strip() for x in re.split(r"\n{2,}|(?=^#{1,4}\s)", text, flags=re.M) if x.strip()]
    out, current = [], ""
    for block in blocks:
        if current and len(current) + len(block) > max_chars:
            out.append(current); current = ""
        current += ("\n\n" if current else "") + block
    if current: out.append(current)
    return out


async def run():
    root = get_settings().course_root
    seen = set()
    async with SessionLocal() as db:
        for dirname, authority in ALLOWED_DIRS.items():
            folder = root / dirname
            if not folder.is_dir(): continue
            for path in folder.rglob("*"):
                if not path.is_file() or path.suffix.lower() not in EXTENSIONS: continue
                relative = str(path.relative_to(root)); seen.add(relative)
                await db.execute(delete(SourceChunk).where(SourceChunk.source_path == relative))
                for locator, text in extract(path):
                    for index, part in enumerate(chunks(text)):
                        digest = hashlib.sha256(f"{relative}:{locator}:{index}:{part}".encode()).hexdigest()
                        db.add(SourceChunk(chunk_id=digest, source_path=relative, source_type=path.suffix[1:],
                            title=path.stem, locator=locator, visibility="student", status="active",
                            authority_level=authority, text=part, content_hash=hashlib.sha256(part.encode()).hexdigest()))
        await db.execute(delete(SourceChunk).where(SourceChunk.source_path.not_in(seen)))
        await db.commit()


if __name__ == "__main__": asyncio.run(run())
