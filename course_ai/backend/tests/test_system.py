"""Local API regression tests; no NAS data or DeepSeek calls are used."""

import json
import io
import time
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient


class CourseAISystemTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.sandbox = tempfile.TemporaryDirectory(prefix="course-ai-tests-")
        root = Path(cls.sandbox.name)
        os.environ.update({
            "DATABASE_URL": "sqlite+aiosqlite:///" + str(root / "test.db"),
            "DEEPSEEK_API_KEY": "CHANGE_ME",
            "DEVICE_TOKEN_SECRET": "local-test-only-not-a-real-secret",
            "FRONTEND_ORIGIN": "https://wenzi-zhuang.github.io",
            "PUBLIC_API_BASE": "https://example.invalid",
            "ASSISTANT_LIBRARY_ROOT": str(root / "library"),
            "RUNTIME_CONFIG_ROOT": str(root / "config"),
            "SETUP_CODE": "local-test-setup-code",
        })
        from course_ai.backend.app import main
        from course_ai.backend.app import case_study_routes
        from course_ai.backend.app.services import case_worker
        cls.main = main

        async def fake_stream(_messages, **_kwargs):
            yield "测试答复"

        main.stream_messages = fake_stream
        case_study_routes.stream_messages = fake_stream
        case_worker.stream_messages = fake_stream
        cls.client = TestClient(main.app)
        cls.client.__enter__()

    @classmethod
    def tearDownClass(cls):
        cls.client.__exit__(None, None, None)
        cls.sandbox.cleanup()

    def credentials(self, student_no="2011310411", suffix="123456", assistant="financial_accounting", **extra):
        return {"student_no": student_no, "id_last6": suffix, "assistant_id": assistant, **extra}

    def test_01_registration_login_and_admin(self):
        self.assertEqual(self.client.post("/api/auth/login", json=self.credentials()).status_code, 403)
        preview = self.client.post("/api/auth/preview", json=self.credentials())
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.json()["name"], "邓伦")
        self.assertEqual(preview.json()["class_name"], "待填写")
        self.assertEqual(self.client.post("/api/auth/register", json=self.credentials()).status_code, 200)
        self.assertEqual(self.client.post("/api/auth/register", json=self.credentials()).status_code, 409)
        login = self.client.post("/api/auth/login", json=self.credentials())
        self.assertEqual(login.status_code, 200)
        self.__class__.token = login.json()["token"]
        self.assertEqual(self.client.get("/api/auth/me?assistant_id=financial_accounting",
                                         headers=self.headers()).json()["name"], "邓伦")
        setup_page = self.client.get("/setup/admin")
        self.assertEqual(setup_page.status_code, 200)
        self.assertEqual(setup_page.text.count("name='password'"), 1)
        self.assertNotIn("minlength=", setup_page.text)
        self.assertEqual(self.client.post("/setup/admin", data={"setup_code": "wrong", "password": "very-strong-password"}).status_code, 403)
        self.assertEqual(self.client.post("/setup/admin", data={"setup_code": "local-test-setup-code", "password": "140210"}).status_code, 200)
        self.assertEqual(self.client.post("/setup/admin", data={"setup_code": "local-test-setup-code", "password": "very-strong-password"}).status_code, 200)
        teacher = self.client.post("/api/auth/login", json=self.credentials(student_no="admin", suffix="very-strong-password"))
        self.assertEqual(teacher.status_code, 200)
        roster = self.client.get("/api/admin/roster?assistant_id=financial_accounting",
                                 headers={"Authorization": "Bearer " + teacher.json()["token"]})
        self.assertEqual(roster.status_code, 200)
        self.assertEqual(len(roster.json()["students"]), 2)
        self.assertEqual(self.client.get("/api/admin/roster?assistant_id=empirical_methods",
                                         headers={"Authorization": "Bearer " + teacher.json()["token"]}).status_code, 403)
        self.assertEqual(self.client.post("/api/auth/login", json=self.credentials(
            student_no="admin", suffix="very-strong-password", assistant="empirical_methods")).status_code, 200)
        self.assertEqual(self.client.post("/setup/admin", data={"setup_code": "local-test-setup-code", "password": "new"}).status_code, 200)
        self.assertEqual(self.client.get("/api/auth/me?assistant_id=financial_accounting",
                                         headers={"Authorization": "Bearer " + teacher.json()["token"]}).status_code, 401)
        self.assertEqual(self.client.post("/api/auth/login", json=self.credentials(
            student_no="admin", suffix="new")).status_code, 200)

    def headers(self):
        return {"Authorization": "Bearer " + self.token}

    def test_02_authz_threads_and_device_cap(self):
        base = "/api/assistants/financial_accounting/threads"
        self.assertEqual(self.client.get(base).status_code, 401)
        body = {"title": "第一问", "messages": [{"role": "user", "content": "第一问"},
                                           {"role": "assistant", "content": "回答"}]}
        self.assertEqual(self.client.put(base + "/" + "a" * 32, json=body, headers=self.headers()).status_code, 200)
        self.assertEqual(len(self.client.get(base, headers=self.headers()).json()["threads"]), 1)
        self.assertEqual(self.client.get(base + "/" + "a" * 32, headers=self.headers()).json()["messages"], body["messages"])
        bad = {"title": "大文件", "messages": [{"role": "user", "content": "x" * 12000}] * 20}
        self.assertEqual(self.client.put(base + "/" + "b" * 32, json=bad, headers=self.headers()).status_code, 413)
        self.assertEqual(self.client.post("/api/auth/register", json=self.credentials(student_no="2013420430", suffix="654321")).status_code, 200)
        other = self.client.post("/api/auth/login", json=self.credentials(student_no="2013420430", suffix="654321"))
        self.assertEqual(other.status_code, 200)
        self.assertEqual(self.client.get(base + "/" + "a" * 32,
                                         headers={"Authorization": "Bearer " + other.json()["token"]}).status_code, 404)
        for _ in range(4):
            self.assertEqual(self.client.post("/api/auth/login", json=self.credentials()).status_code, 200)
        self.assertEqual(self.client.post("/api/auth/login", json=self.credentials()).status_code, 403)
        replacement = self.client.post("/api/auth/login", json=self.credentials(previous_token=self.token))
        self.assertEqual(replacement.status_code, 200)
        self.__class__.token = replacement.json()["token"]
        self.assertEqual(self.client.get("/api/auth/me?assistant_id=financial_accounting",
                                         headers={"Authorization": "Bearer " + self.token}).status_code, 200)

    def test_03_upload_quota_and_protected_chat(self):
        upload = "/api/assistants/financial_accounting/files"
        self.assertEqual(self.client.post(upload, files={"file": ("a.txt", b"hello")}).status_code, 401)
        self.assertEqual(self.client.post(upload, files={"file": ("big.txt", b"x" * (500 * 1024 + 1))},
                                          headers=self.headers()).status_code, 413)
        for _ in range(10):
            response = self.client.post(upload, files={"file": ("a.txt", b"hello")}, headers=self.headers())
            self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.client.post(upload, files={"file": ("a.txt", b"hello")},
                                          headers=self.headers()).status_code, 429)
        chat = "/api/assistants/financial_accounting/chat/stream"
        self.assertEqual(self.client.post(chat, json={"messages": [{"role": "user", "content": "你好"}]}).status_code, 401)
        for _ in range(50):
            response = self.client.post(chat, json={"messages": [{"role": "user", "content": "你好"}]},
                                        headers=self.headers())
            self.assertEqual(response.status_code, 200, response.text)
            self.assertIn("测试答复", response.text)
        self.assertEqual(self.client.post(chat, json={"messages": [{"role": "user", "content": "你好"}]},
                                          headers=self.headers()).status_code, 429)
        self.assertEqual(self.client.post("/api/thesis/check").status_code, 404)

    def test_04_cumulative_bytes_cors_and_device_switch(self):
        empirical = self.credentials(student_no="2013420430", suffix="654321", assistant="empirical_methods")
        self.assertEqual(self.client.post("/api/auth/login", json=empirical).status_code, 403)
        self.assertEqual(self.client.post("/api/auth/preview", json=empirical).status_code, 200)
        self.assertEqual(self.client.post("/api/auth/register", json=empirical).status_code, 200)
        self.assertEqual(self.client.post("/api/auth/register", json=empirical).status_code, 409)
        self.assertEqual(self.client.get("/api/auth/me?assistant_id=empirical_methods",
                                         headers=self.headers()).status_code, 403)
        self.assertEqual(self.client.get("/api/assistants/empirical_methods/threads",
                                         headers=self.headers()).status_code, 403)
        other_login = self.client.post("/api/auth/login", json=self.credentials(
            student_no="2013420430", suffix="654321", assistant="empirical_methods"))
        self.assertEqual(other_login.status_code, 200)
        other_token = other_login.json()["token"]
        url = "/api/assistants/empirical_methods/files"
        for _ in range(3):
            response = self.client.post(url, files={"file": ("large.txt", b"a" * (500 * 1024))},
                                        headers={"Authorization": "Bearer " + other_token})
            self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.client.post(url, files={"file": ("small.txt", b"a")},
                                          headers={"Authorization": "Bearer " + other_token}).status_code, 429)
        self.assertEqual(self.client.get("/api/admin/roster?assistant_id=financial_accounting",
                                         headers=self.headers()).status_code, 403)
        preflight = self.client.options("/api/auth/login", headers={"Origin": "https://evil.invalid",
                                                                 "Access-Control-Request-Method": "POST"})
        self.assertNotIn("https://evil.invalid", preflight.headers.get("access-control-allow-origin", ""))
        switched = self.client.post("/api/auth/login", json=self.credentials(
            student_no="2013420430", suffix="654321", previous_token=self.token))
        self.assertEqual(switched.status_code, 200)
        self.assertEqual(self.client.get("/api/auth/me?assistant_id=financial_accounting",
                                         headers=self.headers()).status_code, 401)
        self.assertEqual(self.client.get("/api/auth/me?assistant_id=empirical_methods",
                                         headers={"Authorization": "Bearer " + other_token}).status_code, 200)

    def test_05_same_student_needs_separate_registration_in_each_course(self):
        another = self.credentials(student_no="2011310411", assistant="case_study")
        self.assertEqual(self.client.post("/api/auth/login", json=another).status_code, 403)
        self.assertEqual(self.client.post("/api/auth/preview", json=another).status_code, 200)
        self.assertEqual(self.client.post("/api/auth/register", json=another).status_code, 200)
        login = self.client.post("/api/auth/login", json=another)
        self.assertEqual(login.status_code, 200)
        token = login.json()["token"]
        self.__class__.case_token = token
        from course_ai.backend.app.services.student_files import student_dir
        device_file = student_dir("case_study", "2011310411") / "devices.json"
        recorded = json.loads(device_file.read_text(encoding="utf-8"))
        self.assertEqual(len(recorded), 1)
        self.assertTrue(recorded[0]["device_id"])
        self.assertEqual(self.client.get("/api/auth/me?assistant_id=case_study",
                                         headers={"Authorization": "Bearer " + token}).status_code, 200)
        self.assertEqual(self.client.get("/api/auth/me?assistant_id=financial_accounting",
                                         headers={"Authorization": "Bearer " + token}).status_code, 403)
        self.assertEqual(self.client.post("/api/assistants/financial_accounting/chat/stream",
                                          json={"messages": [{"role": "user", "content": "你好"}]},
                                          headers={"Authorization": "Bearer " + token}).status_code, 403)
        self.assertEqual(self.client.post("/api/auth/preview", json=self.credentials(
            student_no="2011310411", suffix="000000", assistant="empirical_study")).status_code, 401)

    def test_06_case_paper_checks_bind_to_exact_files(self):
        from docx import Document
        from docx.shared import Cm
        from course_ai.backend.app.services.student_files import student_dir

        def word_bytes(lines):
            document = Document()
            section = document.sections[0]
            section.page_width, section.page_height = Cm(21), Cm(29.7)
            section.top_margin = section.bottom_margin = Cm(2)
            section.left_margin = section.right_margin = Cm(2)
            for line in lines:
                document.add_paragraph(line)
            output = io.BytesIO()
            document.save(output)
            return output.getvalue()

        paper = word_bytes(["东北财经大学本科毕业论文", "案例研究标题", "摘  要", "关键词：案例  会计",
                            "ABSTRACT", "Key words: case; accounting", "目  录", "一、引言",
                            "主要参考文献", "后  记", "指导教师评语"])
        changed = word_bytes(["东北财经大学本科毕业论文", "另一版本", "摘  要"])
        evaluation = word_bytes([f"{index} 这是第 {index} 项的学生回复。" for index in range(12, 19)])
        headers = {"Authorization": "Bearer " + self.case_token}
        base = "/api/case-study"
        self.assertEqual(self.client.post(base + "/format-check", files={"file": ("a.docx", paper)}).status_code, 401)
        premature = self.client.post(base + "/submit", files={
            "file": ("邓伦2011310411清洁版论文.docx", paper),
            "self_eval": ("邓伦2011310411自评表.docx", evaluation)}, headers=headers)
        self.assertEqual(premature.status_code, 409, premature.text)
        format_result = self.client.post(base + "/format-check", files={"file": ("a.docx", paper)}, headers=headers)
        self.assertEqual(format_result.status_code, 200, format_result.text)
        self.assertEqual(len(format_result.json()["items"]), 47)
        self.assertEqual(format_result.json()["items"][0]["status"], "通过")
        content_result = self.client.post(base + "/content-check", files={
            "file": ("a.docx", paper), "self_eval": ("b.docx", evaluation)}, headers=headers)
        self.assertEqual(content_result.status_code, 200, content_result.text)
        self.assertEqual(len(content_result.json()["structure"]), 9)
        altered = self.client.post(base + "/submit", files={
            "file": ("邓伦2011310411清洁版论文.docx", changed),
            "self_eval": ("邓伦2011310411自评表.docx", evaluation)}, headers=headers)
        self.assertEqual(altered.status_code, 409, altered.text)
        bad_name = self.client.post(base + "/submit", files={
            "file": ("wrong.docx", paper), "self_eval": ("邓伦2011310411自评表.docx", evaluation)}, headers=headers)
        self.assertEqual(bad_name.status_code, 400, bad_name.text)
        submitted = self.client.post(base + "/submit", files={
            "file": ("邓伦2011310411清洁版论文.docx", paper),
            "self_eval": ("邓伦2011310411自评表.docx", evaluation)},
            data={"email": "student@example.com"}, headers=headers)
        self.assertEqual(submitted.status_code, 200, submitted.text)
        folder = student_dir("case_study", "2011310411") / "submissions" / submitted.json()["submission_id"]
        self.assertEqual((folder / "邓伦2011310411清洁版论文.docx").read_bytes(), paper)
        self.assertEqual(json.loads((folder / "submission.json").read_text())["email"], "student@example.com")
        for _ in range(40):
            if (folder / "teacher_review_draft.txt").is_file():
                break
            time.sleep(.05)
        self.assertTrue((folder / "teacher_review_draft.txt").is_file())
        status = self.client.get(base + "/submissions/" + submitted.json()["submission_id"], headers=headers)
        self.assertEqual(status.json()["review_status"], "done")
        self.assertEqual(status.json()["mail_status"], "awaiting_configuration")
        self.__class__.case_submission_id = submitted.json()["submission_id"]

    def test_07_mail_setup_and_background_draft(self):
        from course_ai.backend.app.services.mail_drafts import load_mail_config
        from course_ai.backend.app.services.student_files import student_dir

        payload = {"setup_code": "wrong", "address": "teacher@dufe.edu.cn",
                   "host": "mail.dufe.edu.cn", "app_password": "test-client-password"}
        self.assertEqual(self.client.post("/setup/mail", data=payload).status_code, 403)
        self.assertIsNone(load_mail_config())
        payload["setup_code"] = "local-test-setup-code"
        saved_calls = []
        folder = student_dir("case_study", "2011310411") / "submissions" / self.case_submission_id
        def fake_append(config, submission_id, subject, recipient, body):
            saved_calls.append((config.address, submission_id, subject, recipient, len(body)))
        with patch("course_ai.backend.app.mail_setup_routes.discover_drafts", return_value="Drafts"), \
                patch("course_ai.backend.app.services.case_worker.append_review_draft", side_effect=fake_append):
            response = self.client.post("/setup/mail", data=payload)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertNotIn("test-client-password", response.text)
            for _ in range(40):
                if json.loads((folder / "submission.json").read_text()).get("mail_status") == "saved":
                    break
                time.sleep(.05)
        self.assertTrue(saved_calls)
        self.assertEqual(saved_calls[0][1], self.case_submission_id)
        config_path = Path(self.sandbox.name) / "config" / "dufe_mail.json"
        self.assertEqual(config_path.stat().st_mode & 0o777, 0o600)
        self.assertEqual(json.loads((folder / "submission.json").read_text())["mail_status"], "saved")

    def test_08_teacher_initial_draft_uses_nas_worker_without_self_eval(self):
        from docx import Document
        from course_ai.backend.app.services.student_files import student_dir

        document = Document()
        document.add_paragraph("研究方法与论文写作课程研究设计初稿")
        document.add_paragraph("数字化转型可能通过内部控制影响会计信息质量。")
        output = io.BytesIO()
        document.save(output)
        paper = output.getvalue()
        url = "/api/case-study/teacher-initial-draft"
        self.assertEqual(self.client.post(url, files={"file": ("初稿.docx", paper)}).status_code, 401)
        self.assertEqual(self.client.post(url, files={"file": ("初稿.docx", paper)},
                                          headers={"Authorization": "Bearer " + self.case_token}).status_code, 403)
        teacher = self.client.post("/api/auth/login", json=self.credentials(
            student_no="admin", suffix="new", assistant="case_study"))
        self.assertEqual(teacher.status_code, 200, teacher.text)
        headers = {"Authorization": "Bearer " + teacher.json()["token"]}
        calls = []
        def fake_append(config, submission_id, subject, recipient, body):
            calls.append((submission_id, subject, recipient, body))
        with patch("course_ai.backend.app.services.case_worker.append_review_draft", side_effect=fake_append):
            response = self.client.post(url, files={"file": ("研究设计初稿.docx", paper)}, headers=headers)
            self.assertEqual(response.status_code, 200, response.text)
            submission_id = response.json()["submission_id"]
            folder = student_dir("case_study", "admin") / "submissions" / submission_id
            for _ in range(40):
                if json.loads((folder / "submission.json").read_text()).get("mail_status") == "saved":
                    break
                time.sleep(.05)
        self.assertEqual((folder / "研究设计初稿.docx").read_bytes(), paper)
        self.assertEqual(json.loads((folder / "submission.json").read_text())["workflow"],
                         "teacher_initial_draft_test")
        self.assertEqual(len(calls), 1)
        self.assertTrue(calls[0][1].startswith("【NAS自动测试｜初稿意见】"))
        self.assertIsNone(calls[0][2])
        self.assertIn("测试答复", calls[0][3])
        self.assertEqual(self.client.get("/api/case-study/submissions/" + submission_id,
                                         headers=headers).json()["mail_status"], "saved")


if __name__ == "__main__":
    unittest.main()
